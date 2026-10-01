import io, json, re, sqlite3, zipfile, hashlib
from difflib import SequenceMatcher, get_close_matches
from datetime import datetime, timezone
from pathlib import Path
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

st.set_page_config(page_title="Data Vault Studio", page_icon="🧱", layout="wide")
DB=Path(__file__).with_name("data_vault_studio.db")
OK={"Reviewed","Approved"}
D={"dialect":"SQL Server","raw_schema":"RAW_VAULT","business_schema":"BUSINESS_VAULT","stage_schema":"STAGE","hub_prefix":"HUB_","link_prefix":"LINK_","sat_prefix":"SAT_","hash_suffix":"_HK","hash_type":"BINARY(32)","timestamp_type":"DATETIME2(7)","record_source_type":"VARCHAR(200)","load_dts":"LOAD_DTS","load_end_dts":"LOAD_END_DTS","record_source":"RECORD_SOURCE","hashdiff":"HASHDIFF","delimiter":"||","null_token":"^^","business_pattern":"^[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)*$","similarity_threshold":0.78,"block_release":True}

def n(v): return re.sub(r"_+","_",re.sub(r"[^A-Za-z0-9]+","_",str(v).strip())).strip("_").upper()
def y(v): return str(v).strip().lower() in {"y","yes","true","1","bk"}
def cl(df):
    x=df.copy(); x.columns=[n(c).lower() for c in x.columns]; return x.fillna("")
def cv(r,k,d=""): return r[k] if k in r.index else d
def parts(v): return [x.strip() for x in str(v).split(" -> ") if x.strip()]
def q(v,d): return f"[{v}]" if d=="SQL Server" else f'"{v}"'
def esc(v): return str(v).replace('"','\\"')

def init_db():
    with sqlite3.connect(DB) as c:
        c.execute("CREATE TABLE IF NOT EXISTS versions(project TEXT,version INTEGER,created_at TEXT,created_by TEXT,comment TEXT,payload TEXT,PRIMARY KEY(project,version))")
        c.execute("CREATE TABLE IF NOT EXISTS naming_repository(layer TEXT,table_name TEXT,column_name TEXT,business_term TEXT,variant TEXT,source TEXT,created_at TEXT,PRIMARY KEY(layer,table_name,column_name,variant))")
def save_version(project,user,comment,payload):
    with sqlite3.connect(DB) as c:
        v=c.execute("SELECT COALESCE(MAX(version),0)+1 FROM versions WHERE project=?",(project,)).fetchone()[0]; c.execute("INSERT INTO versions VALUES(?,?,?,?,?,?)",(project,v,datetime.now(timezone.utc).isoformat(),user,comment,json.dumps(payload,default=str)))
    return v
def history(project):
    with sqlite3.connect(DB) as c: return pd.read_sql_query("SELECT version,created_at,created_by,comment FROM versions WHERE project=? ORDER BY version DESC",c,params=(project,))
def restore(project,v):
    with sqlite3.connect(DB) as c: r=c.execute("SELECT payload FROM versions WHERE project=? AND version=?",(project,int(v))).fetchone()
    return json.loads(r[0]) if r else None

def template():
    wb=Workbook(); ws=wb.active; ws.title="README"; ws.sheet_state="visible"
    for r in [["Data Vault Studio metadata template"],["1","Populate metadata, one row per source column"],["2","Flag business keys with Y"],["3","Add relationships for Link candidates"],["4","Populate existing_model and naming_dictionary for standards checks"],["5","Optional flow_nodes and flow_edges enrich the data-flow diagram"]]: ws.append(r)
    sheets={
      "metadata":(["table_name","column_name","data_type","is_business_key","source_system","description","rate_of_change","security_class","is_multi_active_key","is_effectivity_attribute","is_status_attribute","satellite_group"],[
       ["CUSTOMER","CORP_CODE","VARCHAR(20)","Y","CRM","Corporation code","LOW","STANDARD","N","N","N",""],["CUSTOMER","CUSTOMER_NAME","VARCHAR(200)","N","CRM","Name","LOW","STANDARD","N","N","N","CORE"],["CUSTOMER","EMAIL","VARCHAR(320)","N","CRM","Email","MEDIUM","PII","N","N","N","PII"],["ACCOUNT","ACCOUNT_ID","VARCHAR(50)","Y","CORE","Account id","LOW","STANDARD","N","N","N",""],["ACCOUNT","CORP_CODE","VARCHAR(20)","N","CORE","Corporation","LOW","STANDARD","N","N","N","CORE"]]),
      "relationships":(["parent_table","parent_column","child_table","child_column","relationship_name","cardinality","link_type","is_transactional","is_hierarchical","requires_effectivity_sat","driving_key"],[["CUSTOMER","CORP_CODE","ACCOUNT","CORP_CODE","CUSTOMER_ACCOUNT","1:M","STANDARD","N","N","Y","CORP_CODE"]]),
      "existing_model":(["layer","table_name","column_name","business_term","approved"],[["BUSINESS_VAULT","BV_CUSTOMER","CORPORATION_CODE","Corporation Code","Y"],["RAW_VAULT","SAT_CUSTOMER","CORP_CODE","Corporation Code","Y"]]),
      "naming_dictionary":(["variant","canonical_name","business_term","approved"],[["CORP_CODE","CORPORATION_CODE","Corporation Code","Y"],["CORP_CD","CORPORATION_CODE","Corporation Code","Y"],["CORPORATION_CD","CORPORATION_CODE","Corporation Code","Y"]]),
      "flow_nodes":(["node_id","label","layer","node_type"],[["SRC_CRM","CRM","SOURCE","SYSTEM"],["BV_CUSTOMER","BV Customer","BUSINESS_VAULT","TABLE"],["DIM_CUSTOMER","Dim Customer","MART","TABLE"],["POWER_BI","Power BI","CONSUMPTION","REPORTING"]]),
      "flow_edges":(["from_node","to_node","label"],[["BV_CUSTOMER","DIM_CUSTOMER","Publishes"],["DIM_CUSTOMER","POWER_BI","Serves"]]),
      "configuration":(["setting","value"],[[k,v] for k,v in D.items()])}
    for name,(heads,rows) in sheets.items():
        sh=wb.create_sheet(name); sh.append(heads)
        for r in rows: sh.append(r)
    fill=PatternFill("solid",fgColor="1F4E78")
    for sh in wb.worksheets:
        sh.freeze_panes="A2" if sh.title!="README" else None
        for c in sh[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=fill
        for cells in sh.columns: sh.column_dimensions[cells[0].column_letter].width=min(max(len(str(c.value or "")) for c in cells)+2,44)
    out=io.BytesIO(); wb.save(out); return out.getvalue()
@st.cache_data(show_spinner=False)
def get_template(): return template()

def read(upload):
    xl=pd.ExcelFile(upload,engine="openpyxl"); sm={s.lower():s for s in xl.sheet_names}
    if "metadata" not in sm: raise ValueError("metadata sheet is required")
    def rd(name,cols): return cl(pd.read_excel(xl,sm[name])) if name in sm else pd.DataFrame(columns=cols)
    md=rd("metadata",[]); rel=rd("relationships",[]); ex=rd("existing_model",["layer","table_name","column_name","business_term","approved"]); dic=rd("naming_dictionary",["variant","canonical_name","business_term","approved"]); fn=rd("flow_nodes",["node_id","label","layer","node_type"]); fe=rd("flow_edges",["from_node","to_node","label"]); cf=rd("configuration",["setting","value"])
    req={"table_name","column_name","data_type","is_business_key"}
    if not req.issubset(md.columns): raise ValueError("metadata missing: "+", ".join(sorted(req-set(md.columns))))
    return md,rel,ex,dic,fn,fe,dict(zip(cf.setting.astype(str),cf.value)) if not cf.empty else {}

def validate(md,rel):
    z=[]
    for c in ["table_name","column_name","data_type"]:
        if md[c].astype(str).str.strip().eq("").any(): z.append(f"Blank {c}")
    if md.duplicated(["table_name","column_name"]).any(): z.append("Duplicate metadata table/column rows")
    for t,g in md.groupby("table_name"):
        if not g.is_business_key.map(y).any(): z.append(f"{t}: no business key")
    known=set(md.table_name.astype(str))
    for _,r in rel.iterrows():
        if str(cv(r,"parent_table")) not in known or str(cv(r,"child_table")) not in known: z.append(f"Unknown relationship endpoint: {cv(r,'parent_table')} -> {cv(r,'child_table')}")
    return sorted(set(z))

def generate(md,rel,cfg):
    A=[]; M=[]
    for table,g in md.groupby("table_name",sort=False):
        bk=g[g.is_business_key.map(y)]; non=g[~g.is_business_key.map(y)]; tn=n(table); src=str(cv(g.iloc[0],"source_system","SOURCE") or "SOURCE")
        if bk.empty: continue
        hub=cfg["hub_prefix"]+tn; A.append([True,"Candidate","HUB",hub,table,"","STANDARD",src,"",""])
        for _,r in bk.iterrows(): M.append([hub,"BUSINESS_KEY",table,r.column_name,r.column_name,r.data_type,"RAW_VAULT","Candidate",""])
        groups={}
        for _,r in non.iterrows():
            subtype="MULTI_ACTIVE" if y(cv(r,"is_multi_active_key")) else "EFFECTIVITY" if y(cv(r,"is_effectivity_attribute")) else "STATUS_TRACKING" if y(cv(r,"is_status_attribute")) else "STANDARD"
            key=str(cv(r,"satellite_group","")).strip() or f"{n(cv(r,'rate_of_change','STANDARD') or 'STANDARD')}_{n(cv(r,'security_class','STANDARD') or 'STANDARD')}_{subtype}"; groups.setdefault(key,[]).append(r)
        for key,rows in groups.items():
            subtype="MULTI_ACTIVE" if any(y(cv(r,"is_multi_active_key")) for r in rows) else "EFFECTIVITY" if any(y(cv(r,"is_effectivity_attribute")) for r in rows) else "STATUS_TRACKING" if any(y(cv(r,"is_status_attribute")) for r in rows) else "STANDARD"
            sat=cfg["sat_prefix"]+tn+("_"+n(key) if len(groups)>1 or n(key)!="STANDARD_STANDARD_STANDARD" else ""); A.append([True,"Candidate","SAT",sat,table,hub,subtype,src,"",""])
            for r in rows: M.append([sat,"MULTI_ACTIVE_KEY" if y(cv(r,"is_multi_active_key")) else "ATTRIBUTE",table,r.column_name,r.column_name,r.data_type,"RAW_VAULT","Candidate",""])
    for _,r in rel.iterrows():
        pn,cn=n(r.parent_table),n(r.child_table); rn=n(cv(r,"relationship_name","") or f"{pn}_{cn}"); subtype="TRANSACTIONAL" if y(cv(r,"is_transactional")) else "HIERARCHICAL" if y(cv(r,"is_hierarchical")) else n(cv(r,"link_type","STANDARD") or "STANDARD")
        link=cfg["link_prefix"]+rn; A.append([True,"Candidate","LINK",link,f"{r.parent_table} -> {r.child_table}",f"{cfg['hub_prefix']+pn}, {cfg['hub_prefix']+cn}",subtype,"SOURCE","",str(cv(r,"driving_key",""))])
        M += [[link,"PARENT_HUB_KEY",r.parent_table,r.parent_column,cfg["hub_prefix"]+pn+cfg["hash_suffix"],"HASH_KEY","RAW_VAULT","Candidate",""] ,[link,"CHILD_HUB_KEY",r.child_table,r.child_column,cfg["hub_prefix"]+cn+cfg["hash_suffix"],"HASH_KEY","RAW_VAULT","Candidate",""]]
        if y(cv(r,"requires_effectivity_sat")):
            sat=cfg["sat_prefix"]+rn+"_EFFECTIVITY"; A.append([True,"Candidate","SAT",sat,f"{r.parent_table} -> {r.child_table}",link,"EFFECTIVITY","SOURCE","","Generated from relationship"])
            for c,dt in [("EFFECTIVE_FROM_DTS",cfg["timestamp_type"]),("EFFECTIVE_TO_DTS",cfg["timestamp_type"]),("IS_CURRENT","CHAR(1)")]: M.append([sat,"ATTRIBUTE",r.child_table,c,c,dt,"RAW_VAULT","Candidate",""])
    return pd.DataFrame(A,columns=["include","status","artifact_type","artifact_name","source_table","parent_artifact","subtype","source_system","reviewer","review_comment"]),pd.DataFrame(M,columns=["artifact_name","role","source_table","source_column","target_column","data_type","layer","column_status","comment"])

def quality(md,a,m,ex,dic,cfg):
    T=[]; E=[]; S=[]; reviewed=set(a[(a.include==True)&a.status.isin(OK)].artifact_name.astype(str)); scope=m[m.artifact_name.astype(str).isin(reviewed)]
    dup=scope[scope.duplicated(["artifact_name","role","source_table","source_column"],False)]; T.append(["Duplicate source mapping","PASS" if dup.empty else "FAIL",len(dup)])
    for _,r in dup.iterrows(): E.append(["ERROR","DUPLICATE_SOURCE",r.artifact_name,r.source_table,r.source_column,r.target_column,"Remove duplicate"])
    dupt=scope[scope.duplicated(["artifact_name","target_column"],False)]; T.append(["Duplicate target mapping","PASS" if dupt.empty else "FAIL",len(dupt)])
    for _,r in dupt.iterrows(): E.append(["ERROR","DUPLICATE_TARGET",r.artifact_name,r.source_table,r.source_column,r.target_column,"Use unique target"])
    base=scope[scope.role.isin(["BUSINESS_KEY","ATTRIBUTE","MULTI_ACTIVE_KEY"])]; expected=set(zip(md.table_name.astype(str),md.column_name.astype(str))); actual=set(zip(base.source_table.astype(str),base.source_column.astype(str))); missing=expected-actual; pct=100 if not expected else round(100*len(expected&actual)/len(expected),2); T.append(["Source coverage","PASS" if pct==100 else "FAIL",len(missing)])
    for t,c in sorted(missing): E.append(["ERROR","MISSING_COVERAGE","",t,c,"","Map to reviewed Hub/Satellite"])
    raw=scope[(scope.layer.astype(str).str.upper()=="RAW_VAULT")&scope.role.isin(["BUSINESS_KEY","ATTRIBUTE","MULTI_ACTIVE_KEY"])]; bad=raw[raw.apply(lambda r:n(r.source_column)!=n(r.target_column),axis=1)]; T.append(["Raw Vault source naming","PASS" if bad.empty else "FAIL",len(bad)])
    for _,r in bad.iterrows(): E.append(["ERROR","RAW_NAME_CHANGED",r.artifact_name,r.source_table,r.source_column,r.target_column,"Retain source name"])
    badobj=a[~a.artifact_name.astype(str).map(lambda x:bool(re.match(r"^(HUB|LINK|SAT)_[A-Z0-9_]+$",n(x))))]; T.append(["Object naming","PASS" if badobj.empty else "FAIL",len(badobj)])
    ref=set(ex[ex.approved.astype(str).map(y)].column_name.astype(str).map(n)) if not ex.empty else set(); known=list(ref|(set(dic.canonical_name.astype(str).map(n)) if not dic.empty else set()))
    bv=scope[scope.layer.astype(str).str.upper()=="BUSINESS_VAULT"]
    for _,r in bv.iterrows():
        cur=n(r.target_column); exact=dic[dic.variant.astype(str).map(n)==cur] if not dic.empty else pd.DataFrame(); proposed=n(exact.iloc[0].canonical_name) if not exact.empty else (get_close_matches(cur,known,n=1,cutoff=float(cfg["similarity_threshold"])) or [cur])[0]; status="PASS" if cur in ref or cur==proposed else "REVIEW"; S.append([status,r.artifact_name,r.source_table,r.source_column,cur,proposed]);
        if status=="REVIEW": E.append(["WARNING","BUSINESS_NAME_VARIANT",r.artifact_name,r.source_table,r.source_column,cur,f"Consider {proposed}"])
    T.append(["Existing-model naming","PASS" if not S or all(x[0]=="PASS" for x in S) else "REVIEW",sum(x[0]!="PASS" for x in S)])
    return pd.DataFrame(T,columns=["test","status","exceptions"]),pd.DataFrame(E,columns=["severity","rule","artifact","source_table","source_column","target_column","recommendation"]),pd.DataFrame(S,columns=["status","artifact","source_table","source_column","current_name","suggested_name"]),pct

def high_dot(a,m):
    active=a[(a.include==True)&a.status.isin(OK)]; hubs=active[active.artifact_type=="HUB"]; links=active[active.artifact_type=="LINK"]
    z=['digraph G {','rankdir=LR;','graph [bgcolor="white",pad="0.3"];','node [shape=box,style="rounded,filled",fontname="Arial",fillcolor="#DCEEFF",color="#2878B5"];']
    for _,r in hubs.iterrows(): z.append(f'"{esc(r.artifact_name)}" [label="{esc(r.artifact_name)}"];')
    for _,r in links.iterrows():
        z.append(f'"{esc(r.artifact_name)}" [shape=diamond,fillcolor="#FFE2B7",color="#D47B00",label="{esc(r.artifact_name)}\\n{esc(r.subtype)}"];')
        for p in str(r.parent_artifact).split(","): z.append(f'"{esc(p.strip())}" -> "{esc(r.artifact_name)}" [dir=none];')
    return '\n'.join(z+['}'])
def low_dot(a,m,cfg,tech=True,source=False,hashdiff=True):
    active=a[(a.include==True)&a.status.isin(OK)]; z=['digraph G {','rankdir=LR;','graph [bgcolor="white",pad="0.3",nodesep="0.4"];','node [shape=plain,fontname="Arial"];']
    colors={"HUB":"#DCEEFF","LINK":"#FFE2B7","SAT":"#DFF3E4"}
    for _,r in active.iterrows():
        mm=m[m.artifact_name==r.artifact_name]; lines=[]
        if r.artifact_type in ["HUB","LINK"]: lines.append(n(r.artifact_name)+cfg["hash_suffix"]+" : "+cfg["hash_type"])
        elif hashdiff: lines.append(cfg["hashdiff"]+" : "+cfg["hash_type"])
        for _,x in mm.iterrows(): lines.append(f"{n(x.target_column)} : {x.data_type}"+(f" [{x.source_table}.{x.source_column}]" if source else ""))
        if tech: lines += [cfg["load_dts"]+" : "+cfg["timestamp_type"],cfg["record_source"]+" : "+cfg["record_source_type"]]
        rows=''.join(f'<TR><TD ALIGN="LEFT">{esc(v)}</TD></TR>' for v in lines)
        z.append(f'"{esc(r.artifact_name)}" [label=<<TABLE BORDER="1" CELLBORDER="0" CELLSPACING="0"><TR><TD BGCOLOR="{colors.get(r.artifact_type,"#EEEEEE")}"><B>{esc(r.artifact_name)}</B><BR/>{esc(r.subtype)}</TD></TR>{rows}</TABLE>>];')
    for _,r in active[active.artifact_type=="SAT"].iterrows(): z.append(f'"{esc(r.parent_artifact)}" -> "{esc(r.artifact_name)}" [label="has"];')
    for _,r in active[active.artifact_type=="LINK"].iterrows():
        for p in str(r.parent_artifact).split(","): z.append(f'"{esc(p.strip())}" -> "{esc(r.artifact_name)}" [dir=none];')
    return '\n'.join(z+['}'])
def flow_dot(md,a,fn,fe,cfg):
    active=a[(a.include==True)&a.status.isin(OK)]; z=['digraph G {','rankdir=LR;','graph [bgcolor="white",pad="0.3",ranksep="0.7"];','node [shape=box,style="rounded,filled",fontname="Arial"];']
    layers={"SOURCE":"#F3F4F6","STAGE":"#E0F2FE","RAW_VAULT":"#DBEAFE","BUSINESS_VAULT":"#EDE9FE","MART":"#DCFCE7","CONSUMPTION":"#FEF3C7"}
    sources=sorted(set(md.get("source_system",pd.Series(["SOURCE"])).astype(str)))
    for s in sources: z.append(f'"SRC_{n(s)}" [label="{esc(s)}",fillcolor="{layers["SOURCE"]}"];'); z.append(f'"STG_{n(s)}" [label="Stage: {esc(s)}",fillcolor="{layers["STAGE"]}"];'); z.append(f'"SRC_{n(s)}" -> "STG_{n(s)}";')
    for _,r in active.iterrows(): z.append(f'"{esc(r.artifact_name)}" [fillcolor="{layers["RAW_VAULT"]}",label="{esc(r.artifact_name)}"];')
    for _,r in active.iterrows():
        src=n(r.source_system if r.source_system!="SOURCE" else (sources[0] if sources else "SOURCE")); z.append(f'"STG_{src}" -> "{esc(r.artifact_name)}";')
    if not fn.empty:
        for _,r in fn.iterrows(): z.append(f'"{esc(r.node_id)}" [label="{esc(r.label)}",fillcolor="{layers.get(n(r.layer),"#EEEEEE")}"];')
    if not fe.empty:
        for _,r in fe.iterrows(): z.append(f'"{esc(r.from_node)}" -> "{esc(r.to_node)}" [label="{esc(r.label)}"];')
    return '\n'.join(z+['}'])
def mermaid_from_dot(dot): return "%% Graphviz DOT is the canonical export for this diagram\n%% Use the .dot file for editing/rendering\n"+dot



def dbml_type(value):
    """Return a DBML-compatible datatype without changing SQL type semantics."""
    value=str(value or "VARCHAR(255)").strip() or "VARCHAR(255)"
    return f'"{value}"' if " " in value and "(" not in value else value

def dbml_note(value):
    return str(value or "").replace("\\","\\\\").replace("'","\\'").replace("\n"," ")

def export_scope(artifacts,include_candidates=False):
    scoped=artifacts.copy()
    if "include" in scoped.columns:
        scoped=scoped[scoped["include"].fillna(False).astype(bool)]
    if not include_candidates and "status" in scoped.columns:
        scoped=scoped[scoped["status"].astype(str).isin(OK)]
    return scoped

def export_columns(artifact,mappings,cfg,include_technical=True):
    """Build a normalized physical-column list used by both DBML and JSON."""
    kind=n(artifact["artifact_type"]); name=n(artifact["artifact_name"]); rows=[]
    def add(name_,dtype_,role_,pk=False,nullable=False,source_table="",source_column="",comment=""):
        rows.append({"name":n(name_),"data_type":str(dtype_),"role":role_,"primary_key":pk,"nullable":nullable,"source_table":str(source_table or ""),"source_column":str(source_column or ""),"comment":str(comment or "")})
    if kind in {"HUB","LINK"}:
        add(name+cfg["hash_suffix"],cfg["hash_type"],"HASH_KEY",pk=True)
    elif kind=="SAT":
        parent=n(artifact.get("parent_artifact","PARENT"))
        add(parent+cfg["hash_suffix"],cfg["hash_type"],"PARENT_HASH_KEY",pk=True)
        add(cfg["hashdiff"],cfg["hash_type"],"HASHDIFF")
    mm=mappings[mappings["artifact_name"].astype(str)==str(artifact["artifact_name"])]
    for _,r in mm.iterrows():
        role=n(r.get("role","ATTRIBUTE")); dtype_=cfg["hash_type"] if role in {"PARENT_HUB_KEY","CHILD_HUB_KEY"} else r.get("data_type","VARCHAR(255)")
        add(r.get("target_column",r.get("source_column","COLUMN")),dtype_,role,nullable=(kind=="SAT" and role!="MULTI_ACTIVE_KEY"),source_table=r.get("source_table",""),source_column=r.get("source_column",""),comment=r.get("comment",""))
    if include_technical:
        add(cfg["load_dts"],cfg["timestamp_type"],"LOAD_TIMESTAMP",pk=(kind=="SAT"))
        if kind=="SAT" and n(artifact.get("subtype","STANDARD")) in {"STANDARD","MULTI_ACTIVE","STATUS_TRACKING"}:
            add(cfg.get("load_end_dts","LOAD_END_DTS"),cfg["timestamp_type"],"LOAD_END_TIMESTAMP",nullable=True)
        add(cfg["record_source"],cfg["record_source_type"],"RECORD_SOURCE")
    unique={}
    for row in rows: unique.setdefault(row["name"],row)
    return list(unique.values())

def export_relationships(artifacts,mappings,cfg):
    relationships=[]; names=set(artifacts["artifact_name"].astype(str))
    for _,artifact in artifacts.iterrows():
        child=str(artifact["artifact_name"]); kind=n(artifact["artifact_type"])
        if kind=="SAT":
            parent=str(artifact.get("parent_artifact","")).strip()
            if parent and parent in names:
                relationships.append({"relationship_type":"SATELLITE_PARENT","from_table":child,"from_column":n(parent)+cfg["hash_suffix"],"to_table":parent,"to_column":n(parent)+cfg["hash_suffix"],"cardinality":"many-to-one"})
        elif kind=="LINK":
            parents=[x.strip() for x in str(artifact.get("parent_artifact","")).split(",") if x.strip()]
            mm=mappings[(mappings["artifact_name"].astype(str)==child)&mappings["role"].astype(str).isin(["PARENT_HUB_KEY","CHILD_HUB_KEY"])]
            for index,(_,mapping) in enumerate(mm.iterrows()):
                parent=parents[index] if index<len(parents) else ""
                if parent and parent in names:
                    relationships.append({"relationship_type":"LINK_PARTICIPANT","from_table":child,"from_column":n(mapping["target_column"]),"to_table":parent,"to_column":n(parent)+cfg["hash_suffix"],"cardinality":"many-to-one"})
    return relationships

def generate_dbml(project,artifacts,mappings,cfg,include_technical=True,include_candidates=False):
    """Generate dbdiagram.io-compatible DBML for the reviewed model."""
    scoped=export_scope(artifacts,include_candidates); schema=n(cfg["raw_schema"]); rels=export_relationships(scoped,mappings,cfg)
    lines=[f"Project {n(project)} {{",f"  database_type: '{dbml_note(cfg['dialect'])}'","  Note: 'Generated by Data Vault Studio'","}",""]
    for _,artifact in scoped.iterrows():
        description=f"{artifact['artifact_type']} {artifact.get('subtype','STANDARD')} | Source: {artifact.get('source_table','')} | Status: {artifact.get('status','')}"
        lines.append(f"Table {schema}.{n(artifact['artifact_name'])} [note: '{dbml_note(description)}'] {{")
        for column in export_columns(artifact,mappings,cfg,include_technical):
            settings=[]
            if column["primary_key"]: settings.append("pk")
            settings.append("null" if column["nullable"] else "not null")
            source_note=f"Role: {column['role']}"
            if column["source_table"] or column["source_column"]: source_note=f"Source: {column['source_table']}.{column['source_column']} | "+source_note
            if column["comment"]: source_note += " | "+column["comment"]
            settings.append(f"note: '{dbml_note(source_note)}'")
            lines.append(f"  {n(column['name'])} {dbml_type(column['data_type'])} [{', '.join(settings)}]")
        lines += ["}",""]
    for rel in rels:
        lines.append(f"Ref: {schema}.{n(rel['from_table'])}.{n(rel['from_column'])} > {schema}.{n(rel['to_table'])}.{n(rel['to_column'])}")
    return "\n".join(lines).rstrip()+"\n"

def generate_json_export(project,artifacts,mappings,cfg,tests,errors,include_technical=True,include_candidates=False):
    """Generate a portable, integration-ready JSON representation of the model."""
    scoped=export_scope(artifacts,include_candidates); names=set(scoped["artifact_name"].astype(str)); items=[]
    for _,artifact in scoped.iterrows():
        items.append({"name":str(artifact["artifact_name"]),"artifact_type":str(artifact["artifact_type"]),"subtype":str(artifact.get("subtype","STANDARD")),"status":str(artifact.get("status","")),"included":bool(artifact.get("include",True)),"source_table":str(artifact.get("source_table","")),"parent_artifact":str(artifact.get("parent_artifact","")),"source_system":str(artifact.get("source_system","")),"reviewer":str(artifact.get("reviewer","")),"review_comment":str(artifact.get("review_comment","")),"columns":export_columns(artifact,mappings,cfg,include_technical)})
    model={"model_format":"DataVaultStudio.Model.v1","generated_at_utc":datetime.now(timezone.utc).isoformat(),"project":{"name":project,"dialect":cfg["dialect"],"raw_schema":cfg["raw_schema"]},"configuration":cfg,"summary":{"artifact_count":len(items),"hub_count":sum(x["artifact_type"].upper()=="HUB" for x in items),"link_count":sum(x["artifact_type"].upper()=="LINK" for x in items),"satellite_count":sum(x["artifact_type"].upper()=="SAT" for x in items)},"artifacts":items,"relationships":export_relationships(scoped,mappings,cfg),"source_to_target_mappings":mappings[mappings["artifact_name"].astype(str).isin(names)].to_dict("records"),"quality":{"tests":tests.to_dict("records"),"exceptions":errors.to_dict("records")}}
    model["summary"]["relationship_count"]=len(model["relationships"])
    return json.dumps(model,indent=2,ensure_ascii=False,default=str)

def ddl(a,m,cfg):
    d=cfg["dialect"]; out=[f"-- Generated {d} Raw Vault DDL"]
    for _,r in a[(a.include==True)&a.status.isin(OK)].iterrows():
        name=n(r.artifact_name); mm=m[m.artifact_name==r.artifact_name]; cols=[]
        if r.artifact_type in ["HUB","LINK"]: cols.append((name+cfg["hash_suffix"],cfg["hash_type"],"NOT NULL"))
        else: cols += [(n(r.parent_artifact)+cfg["hash_suffix"],cfg["hash_type"],"NOT NULL"),(cfg["hashdiff"],cfg["hash_type"],"NOT NULL")]
        for _,x in mm.iterrows(): cols.append((n(x.target_column),cfg["hash_type"] if x.data_type=="HASH_KEY" else x.data_type,"NULL" if r.artifact_type=="SAT" else "NOT NULL"))
        cols += [(cfg["load_dts"],cfg["timestamp_type"],"NOT NULL"),(cfg["record_source"],cfg["record_source_type"],"NOT NULL")]
        out.append(f"CREATE TABLE {q(n(cfg['raw_schema']),d)}.{q(name,d)} (\n"+",\n".join(f"  {q(c,d)} {t} {null}" for c,t,null in cols)+"\n);\n")
    return '\n'.join(out)
def stage(m,cfg):
    out=["-- Staging views"]
    for src in sorted(set(m.source_table.astype(str))):
        cols=list(dict.fromkeys(m[m.source_table.astype(str)==src].source_column.astype(str))); out.append(f"CREATE VIEW {n(cfg['stage_schema'])}.STG_{n(src)} AS SELECT "+", ".join(q(n(c),cfg["dialect"]) for c in cols)+f" FROM {src};")
    return '\n'.join(out)
def loads(a,cfg):
    return '\n'.join(["-- Load-pattern skeletons. Confirm grain, CDC, driving keys and parent lookups."]+[f"-- {n(r.artifact_name)}: INSERT/MERGE pattern for {r.artifact_type} {r.subtype}" for _,r in a[(a.include==True)&a.status.isin(OK)].iterrows()])
def table_summary(md,a,m):
    rows=[]
    for t in md.table_name.drop_duplicates().astype(str):
        x=a[a.source_table.astype(str).map(lambda v:t in parts(v))]; reviewed=set(x[x.status.isin(OK)].artifact_name.astype(str)); total=len(x); done=len(reviewed); ct=md[md.table_name.astype(str)==t].column_name.nunique(); cd=m[(m.artifact_name.isin(reviewed))&(m.source_table.astype(str)==t)].source_column.nunique(); status="Approved" if total and (x.status=="Approved").all() else "Reviewed" if total and done==total else "Partial" if done else "Candidate"; rows.append([False,t,status,total,done,ct,min(cd,ct),round(100*done/total,1) if total else 0])
    return pd.DataFrame(rows,columns=["select","table_name","status","artifacts","reviewed","columns","covered","review_pct"])
def bundle(files):
    out=io.BytesIO()
    with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
        for name,data in files.items(): z.writestr(name,data)
    return out.getvalue()

def main():
    init_db(); st.title("Data Vault Studio"); st.caption("Full modeling, review, quality, ER visualization, data flow, code generation and governance")

    with st.sidebar:
        st.header("Getting started"); st.download_button("Download metadata template",get_template(),"data_vault_studio_template.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True); up=st.file_uploader("Upload completed workbook",type=["xlsx"]); project=st.text_input("Project","RAW_VAULT_ACCELERATOR"); reviewer=st.text_input("Reviewer","Arbind"); dialect=st.selectbox("Target",["SQL Server","Snowflake","PostgreSQL"]); raw=st.text_input("Raw Vault schema","RAW_VAULT"); stg=st.text_input("Stage schema","STAGE")
    if not up: st.info("Download the template, populate it, and upload the completed workbook."); return
    try: md,rel,ex,dic,fn,fe,fc=read(up)
    except Exception as e: st.error(str(e)); return
    cfg=D.copy(); cfg.update({k:v for k,v in fc.items() if k in cfg}); cfg.update(dialect=dialect,raw_schema=raw,stage_schema=stg)
    if dialect=="Snowflake": cfg.update(hash_type="BINARY(32)",timestamp_type="TIMESTAMP_NTZ")
    elif dialect=="PostgreSQL": cfg.update(hash_type="BYTEA",timestamp_type="TIMESTAMP")
    fid=hashlib.sha256(up.getvalue()).hexdigest()
    if st.session_state.get("fid")!=fid: st.session_state.a,st.session_state.m=generate(md,rel,cfg); st.session_state.fid=fid
    a,m=st.session_state.a,st.session_state.m; issues=validate(md,rel)
    names=["Dashboard","Validation","Metadata","Relationships","Candidate model","Table & bulk review","Artifact review","Column mapping","Quality gate","Naming repository","High-level ER","Low-level ER","Data flow diagram","Configuration","DDL","Staging SQL","Load patterns","Versions","Export"]
    tabs=st.tabs(names)
    with tabs[0]:
        ts=table_summary(md,a,m); rev=int(a.status.isin(OK).sum()); rn=set(a[a.status.isin(OK)].artifact_name); covered=len(set(zip(m[m.artifact_name.isin(rn)].source_table,m[m.artifact_name.isin(rn)].source_column))&set(zip(md.table_name,md.column_name))); c1,c2,c3=st.columns(3); c1.metric("Tables reviewed",f"{sum(ts.status.isin(OK))}/{len(ts)}"); c2.metric("Artifacts reviewed",f"{rev}/{len(a)}"); c3.metric("Columns covered",f"{covered}/{len(md)}"); st.dataframe(ts.drop(columns="select"),use_container_width=True,hide_index=True)
    with tabs[1]:
        if issues:
            for i in issues: st.warning(i)
        else: st.success("Input validation passed")
    with tabs[2]: st.dataframe(md,use_container_width=True,hide_index=True)
    with tabs[3]: st.dataframe(rel,use_container_width=True,hide_index=True)
    with tabs[4]: st.dataframe(a,use_container_width=True,hide_index=True)
    with tabs[5]:
        x=st.data_editor(table_summary(md,a,m),use_container_width=True,hide_index=True,disabled=["table_name","status","artifacts","reviewed","columns","covered","review_pct"],key="tables"); chosen=set(x[x.select].table_name.astype(str)); cs=st.columns(6)
        for c,(label,status) in zip(cs,[("Review selected","Reviewed"),("Approve selected","Approved"),("Needs change","Needs change"),("Reject","Rejected"),("Reset","Candidate"),("Review all","Reviewed")]):
            if c.button(label,use_container_width=True):
                mask=a.include==True if label=="Review all" else a.source_table.astype(str).map(lambda v:bool(chosen.intersection(parts(v)))); st.session_state.a.loc[mask,["status","reviewer"]]=[status,reviewer]; st.rerun()
    with tabs[6]: st.session_state.a=st.data_editor(a,use_container_width=True,hide_index=True,num_rows="dynamic",column_config={"status":st.column_config.SelectboxColumn(options=["Candidate","Needs change","Reviewed","Approved","Rejected"]),"artifact_type":st.column_config.SelectboxColumn(options=["HUB","LINK","SAT"])},key="arts")
    with tabs[7]: st.session_state.m=st.data_editor(m,use_container_width=True,hide_index=True,num_rows="dynamic",column_config={"layer":st.column_config.SelectboxColumn(options=["RAW_VAULT","BUSINESS_VAULT"]),"column_status":st.column_config.SelectboxColumn(options=["Candidate","Reviewed","Approved","Rejected"])},key="maps")
    tests,errors,suggest,pct=quality(md,st.session_state.a,st.session_state.m,ex,dic,cfg); blocked=not errors[errors.severity=="ERROR"].empty if not errors.empty else False
    with tabs[8]:
        c1,c2,c3=st.columns(3); c1.metric("Checks passed",f"{sum(tests.status=='PASS')}/{len(tests)}"); c2.metric("Coverage",f"{pct}%"); c3.metric("Blocking errors",sum(errors.severity=="ERROR") if not errors.empty else 0); st.dataframe(tests,use_container_width=True,hide_index=True); st.dataframe(errors,use_container_width=True,hide_index=True); st.dataframe(suggest,use_container_width=True,hide_index=True)
    with tabs[9]: st.dataframe(ex,use_container_width=True,hide_index=True); st.dataframe(dic,use_container_width=True,hide_index=True)
    high=high_dot(st.session_state.a,st.session_state.m)
    with tabs[10]:
        st.caption("Executive view: reviewed Hubs and Links only"); st.graphviz_chart(high,use_container_width=True); st.download_button("Download high-level ER (.dot)",high,"high_level_er.dot")
    with tabs[11]:
        c1,c2,c3=st.columns(3); tech=c1.checkbox("Technical columns",True); src=c2.checkbox("Source lineage",False); hd=c3.checkbox("HashDiff",True); low=low_dot(st.session_state.a,st.session_state.m,cfg,tech,src,hd); st.graphviz_chart(low,use_container_width=True); st.download_button("Download low-level ER (.dot)",low,"low_level_er.dot")
    flow=flow_dot(md,st.session_state.a,fn,fe,cfg)
    with tabs[12]:
        st.caption("Source → Stage → reviewed Raw Vault → optional Business Vault / Mart / Consumption nodes"); st.graphviz_chart(flow,use_container_width=True); st.download_button("Download data-flow diagram (.dot)",flow,"data_flow.dot")
    with tabs[13]: st.json(cfg)
    dsql=ddl(st.session_state.a,st.session_state.m,cfg); ssql=stage(st.session_state.m,cfg); lsql=loads(st.session_state.a,cfg)
    with tabs[14]:
        if blocked and y(cfg["block_release"]): st.error("DDL release blocked by quality errors")
        st.code(dsql,"sql")
    with tabs[15]: st.code(ssql,"sql")
    with tabs[16]: st.code(lsql,"sql")
    with tabs[17]:
        comment=st.text_input("Version comment")
        if st.button("Save version"): st.success(f"Saved version {save_version(project,reviewer,comment,{'configuration':cfg,'artifacts':st.session_state.a.to_dict('records'),'mappings':st.session_state.m.to_dict('records')})}")
        h=history(project); st.dataframe(h,use_container_width=True,hide_index=True)
        if not h.empty:
            v=st.selectbox("Restore version",h.version.tolist())
            if st.button("Restore"): p=restore(project,v); st.session_state.a=pd.DataFrame(p["artifacts"]); st.session_state.m=pd.DataFrame(p["mappings"]); st.rerun()
    with tabs[18]:
        review = io.BytesIO()
        # Correctly pass w as excel_writer, and clearly name each sheet string target
        with pd.ExcelWriter(review, engine="openpyxl") as w:
            pd.DataFrame({"status": ["Data Vault Studio review package"]}).to_excel(excel_writer=w, sheet_name="README", index=False)
            st.session_state.a.to_excel(excel_writer=w, sheet_name="artifacts", index=False)
            st.session_state.m.to_excel(excel_writer=w, sheet_name="mappings", index=False)
            tests.to_excel(excel_writer=w, sheet_name="quality_tests", index=False)
            errors.to_excel(excel_writer=w, sheet_name="quality_exceptions", index=False)
        st.subheader("DBML and JSON export")
        e1,e2=st.columns(2)
        include_technical=e1.checkbox("Include technical columns",True,key="export_technical")
        include_candidates=e2.checkbox("Include Candidate and Needs Change artifacts",False,key="export_candidates")
        dbml_text=generate_dbml(project,st.session_state.a,st.session_state.m,cfg,include_technical,include_candidates)
        json_text=generate_json_export(project,st.session_state.a,st.session_state.m,cfg,tests,errors,include_technical,include_candidates)
        preview_dbml,preview_json=st.tabs(["DBML preview","JSON preview"])
        with preview_dbml: st.code(dbml_text,language="text")
        with preview_json: st.code(json_text,language="json")
        d1,d2,d3=st.columns(3)
        d1.download_button("Download DBML",dbml_text,f"{project}_raw_vault.dbml","text/plain",use_container_width=True)
        d2.download_button("Download JSON",json_text,f"{project}_model.json","application/json",use_container_width=True)
        dbml_json_zip=io.BytesIO()
        with zipfile.ZipFile(dbml_json_zip,"w",zipfile.ZIP_DEFLATED) as z:
            z.writestr(f"{project}_raw_vault.dbml",dbml_text)
            z.writestr(f"{project}_model.json",json_text)
        d3.download_button("Download DBML + JSON",dbml_json_zip.getvalue(),f"{project}_dbml_json.zip","application/zip",use_container_width=True)
        files={"raw_vault_ddl.sql":dsql,"staging_views.sql":ssql,"load_patterns.sql":lsql,"high_level_er.dot":high,"low_level_er.dot":low_dot(st.session_state.a,st.session_state.m,cfg),"data_flow.dot":flow,"high_level_er.mmd":mermaid_from_dot(high),"low_level_er.mmd":mermaid_from_dot(low_dot(st.session_state.a,st.session_state.m,cfg)),"data_flow.mmd":mermaid_from_dot(flow),"review_package.xlsx":review.getvalue(),f"{project}_raw_vault.dbml":dbml_text,f"{project}_model.json":json_text,"quality_tests.csv":tests.to_csv(index=False),"quality_exceptions.csv":errors.to_csv(index=False)}
        if blocked and y(cfg["block_release"]): st.error("Deployment package blocked. Review package remains available.")
        st.download_button("Download complete deployment package",bundle(files),"data_vault_studio_package.zip",disabled=blocked and y(cfg["block_release"])); st.download_button("Download review package",review.getvalue(),"review_package.xlsx")
if __name__=="__main__": main()
